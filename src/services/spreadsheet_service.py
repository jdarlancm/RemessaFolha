"""Service for managing control spreadsheets."""
import pdfplumber

from typing import List, Dict, Optional
from openpyxl import load_workbook
from pypdf import PdfReader

from domain.payroll import PayrollRemittance
from helpers import paths_helper, paycheck_helper
from utils import date_utils

class SpreadsheetService:
  """Service for managing payroll control spreadsheets."""
  
  def update_control_sheet(self, remittance: PayrollRemittance) -> None:
    """
    Update the control spreadsheet with payroll data.
    
    Args:
        remittance: The payroll remittance data
    """
    # Extrair salários dos contracheques
    payroll_salaries = self._extract_payroll_salaries(remittance)
    
    # Atualizar planilha
    self._update_payroll_sheet(
      remittance.reference_date.month,
      remittance.reference_date.year,
      payroll_salaries
    )
    
  def _extract_payroll_salaries(
    self,
    remittance: PayrollRemittance
  ) -> List[Dict[str, str]]:
    """Extract salary information from paychecks."""
    payroll_salaries = []
    
    for paycheck in remittance.paychecks:
      if not paycheck.file_path:
        continue
        
      # Extrair dados do contracheque
      with pdfplumber.open(paycheck.file_path) as pdf:
        page_content = pdf.pages[0].extract_text()

        # Verificar se é um contracheque
        if not "RECIBO DE PAGAMENTO DE SALÁRIO" in page_content:
          continue
        
        # Extrair matrícula e salário
        matricula = paycheck_helper.extract_matricula(page_content)
        salario_liquido = paycheck_helper.extract_net_salary(page_content)
      
        if not salario_liquido or matricula == 0:
          raise ValueError(
            f"c) Não foi possível obter informações do funcionário no contra-cheque "
            f"({paycheck.file_path})"
          )
        
        payroll_salaries.append({
          "matricula": matricula,
          "salario_liquido": salario_liquido
        })
      
    return payroll_salaries
    
  def _update_payroll_sheet(
    self,
    month: int,
    year: int,
    payroll_salaries: List[Dict[str, str]]
  ) -> None:
    """Update the payroll control spreadsheet."""
    COL_MATRICULA = 1
    
    # Abrir planilha
    payroll_filename = paths_helper.get_payroll_year_complete_filename(year)
    spreadsheet_tab = "folha"
    
    wb = load_workbook(filename=payroll_filename)
    spreadsheet = wb[spreadsheet_tab]
    
    # Encontrar coluna do mês
    col_reference_month = self._get_reference_month(month, spreadsheet)
    if not col_reference_month:
      raise ValueError(
        f"Não foi encontrada a coluna de referência do mes {year}-{month} "
        f"na planilha da folha: {payroll_filename}"
      )
      
    # Atualizar valores
    for row in spreadsheet.iter_rows(min_row=3, max_row=spreadsheet.max_row):
      matricula = row[COL_MATRICULA].value
      salario_liquido = self._get_salary_from_payroll(matricula, payroll_salaries)
      if salario_liquido:
        cell_coordinate = f"{col_reference_month}{row[0].row}"
        spreadsheet[cell_coordinate] = float(salario_liquido)
        
    # Salvar planilha
    wb.save(payroll_filename)
    
  def _get_reference_month(
    self,
    month: int,
    spreadsheet
  ) -> Optional[str]:
    """Get the column letter for the reference month."""
    ROW_HEADER = 2
    col_name_reference_month = date_utils.nome_mes(month)
    
    for cell in spreadsheet[ROW_HEADER]:
      if cell.value == col_name_reference_month:
        return cell.column_letter
        
    return None
    
  def _get_salary_from_payroll(
    self,
    matricula: int,
    payroll_data: List[Dict[str, str]]
  ) -> Optional[str]:
    """Get salary for an employee from payroll data."""
    for employee in payroll_data:
      if int(employee["matricula"]) == matricula:
        return employee["salario_liquido"].replace(",", ".")
        
    return None 